from datetime import date
from openpyxl import load_workbook

from personal_finance_agent.item_matcher import match_itemized_purchases
from personal_finance_agent.models import ImportedTransaction
from personal_finance_agent.reports import (
    available_months,
    build_combined_report,
    build_monthly_report,
    category_totals,
    itemized_category_allocations,
    merchant_totals,
)
from personal_finance_agent.storage import connect, insert_itemized_purchases, insert_transactions, transactions_for_month


def test_category_totals_use_scaled_itemized_allocations(tmp_path) -> None:
    conn = connect(tmp_path / "finance.sqlite")
    insert_transactions(
        conn,
        [
            ImportedTransaction(
                source_file="bank.csv",
                source_account="card",
                transaction_date=date(2026, 6, 8),
                raw_description="Target",
                normalized_merchant="Target",
                amount=-20.00,
            )
        ],
    )
    insert_itemized_purchases(
        conn,
        [
            {
                "merchant": "Target",
                "order_id": "T1",
                "order_date": "2026-06-07",
                "item_title": "Groceries",
                "item_price": 10.00,
                "quantity": 1,
                "item_total": 10.00,
                "order_total": 20.00,
                "item_category": "Groceries",
                "confidence": 0.82,
            },
            {
                "merchant": "Target",
                "order_id": "T1",
                "order_date": "2026-06-07",
                "item_title": "Paper towels",
                "item_price": 15.00,
                "quantity": 1,
                "item_total": 15.00,
                "order_total": 20.00,
                "item_category": "Household Consumables",
                "confidence": 0.82,
            },
        ],
    )
    assert match_itemized_purchases(conn, "2026-06") == 1

    allocations = itemized_category_allocations(conn, "2026-06")
    assert {allocation.category: allocation.amount for allocation in allocations} == {
        "Groceries": 8.00,
        "Household & Kids": 12.00,
    }
    assert category_totals(transactions_for_month(conn, "2026-06"), conn, "2026-06") == {
        "Groceries": 8.00,
        "Household & Kids": 12.00,
    }


def test_merchant_totals_split_itemized_transactions_by_category(tmp_path) -> None:
    conn = connect(tmp_path / "finance.sqlite")
    insert_transactions(
        conn,
        [
            ImportedTransaction(
                source_file="bank.csv",
                source_account="card",
                transaction_date=date(2026, 6, 8),
                raw_description="Target",
                normalized_merchant="Target",
                amount=-20.00,
            )
        ],
    )
    insert_itemized_purchases(
        conn,
        [
            {
                "merchant": "Target",
                "order_id": "T1",
                "order_date": "2026-06-07",
                "item_title": "Groceries",
                "item_price": 10.00,
                "quantity": 1,
                "item_total": 10.00,
                "order_total": 20.00,
                "item_category": "Groceries",
                "confidence": 0.82,
            },
            {
                "merchant": "Target",
                "order_id": "T1",
                "order_date": "2026-06-07",
                "item_title": "Paper towels",
                "item_price": 10.00,
                "quantity": 1,
                "item_total": 10.00,
                "order_total": 20.00,
                "item_category": "Household Consumables",
                "confidence": 0.82,
            },
        ],
    )
    assert match_itemized_purchases(conn, "2026-06") == 1

    totals = merchant_totals(transactions_for_month(conn, "2026-06"), conn, "2026-06")

    assert totals[("Target", "Groceries")]["total"] == 10.00
    assert totals[("Target", "Household & Kids")]["total"] == 10.00


def test_build_monthly_report_uses_valid_sheet_titles(tmp_path) -> None:
    conn = connect(tmp_path / "finance.sqlite")
    insert_transactions(
        conn,
        [
            ImportedTransaction(
                source_file="bank.csv",
                source_account="card",
                transaction_date=date(2026, 6, 8),
                raw_description="Target",
                normalized_merchant="Target",
                amount=-20.00,
            )
        ],
    )

    workbook_path, markdown_path = build_monthly_report(conn, "2026-06", tmp_path / "exports")

    workbook = load_workbook(workbook_path)
    assert "Amazon Target Detail" in workbook.sheetnames
    assert markdown_path.exists()


def test_build_combined_report_across_available_months(tmp_path) -> None:
    conn = connect(tmp_path / "finance.sqlite")
    insert_transactions(
        conn,
        [
            ImportedTransaction(
                source_file="bank.csv",
                source_account="card",
                transaction_date=date(2026, 1, 10),
                raw_description="Grocery Store",
                normalized_merchant="Grocery Store",
                amount=-50.00,
            ),
            ImportedTransaction(
                source_file="bank.csv",
                source_account="card",
                transaction_date=date(2026, 2, 10),
                raw_description="Target",
                normalized_merchant="Target",
                amount=-20.00,
            ),
        ],
    )
    conn.execute(
        """
        update transactions
        set category = 'Groceries',
            confidence = 1.0,
            needs_review = 0,
            categorization_source = 'test'
        where normalized_merchant = 'Grocery Store'
        """
    )
    conn.execute(
        """
        update transactions
        set category = 'Household & Kids',
            confidence = 1.0,
            needs_review = 0,
            categorization_source = 'test'
        where normalized_merchant = 'Target'
        """
    )
    conn.commit()
    insert_itemized_purchases(
        conn,
        [
            {
                "merchant": "Target",
                "order_id": "T1",
                "order_date": "2026-02-09",
                "item_title": "Paper towels",
                "item_price": 20.00,
                "quantity": 1,
                "item_total": 20.00,
                "order_total": 20.00,
                "item_category": "Household Consumables",
                "item_subcategory": "Paper Goods",
                "confidence": 0.82,
            }
        ],
    )
    assert match_itemized_purchases(conn, "2026-02") == 1

    workbook_path = build_combined_report(conn, tmp_path / "exports")

    assert workbook_path.name == "combined_review_all.xlsx"
    assert available_months(conn) == ["2026-01", "2026-02"]
    workbook = load_workbook(workbook_path)
    assert {
        "Overview",
        "Monthly Category Trend",
        "Focus Categories",
        "Top Merchants",
        "Recurring Charges",
        "Category Drift",
        "Review Health",
        "Itemized Allocations",
        "Subcategory Detail",
        "Raw Combined Transactions",
    }.issubset(set(workbook.sheetnames))
    trend_headers = [cell.value for cell in workbook["Monthly Category Trend"][1]]
    assert "2026-01" in trend_headers
    assert "2026-02" in trend_headers
    subcategory_rows = list(workbook["Subcategory Detail"].iter_rows(values_only=True))
    assert ("2026-02", "Household & Kids", "Paper Goods", "Target", 1, 20) in subcategory_rows
