from __future__ import annotations

import sqlite3
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from personal_finance_agent.categories import FOCUS_CATEGORIES
from personal_finance_agent.item_enricher import ITEM_TO_BUDGET
from personal_finance_agent.storage import itemized_purchases, seed_default_budgets, transactions_for_month


@dataclass(frozen=True)
class ItemizedAllocation:
    transaction_id: int
    merchant: str
    category: str
    amount: float


@dataclass(frozen=True)
class ItemizedSubcategoryAllocation:
    transaction_id: int
    merchant: str
    category: str
    subcategory: str
    amount: float


def build_monthly_report(conn: sqlite3.Connection, month: str, exports_path: Path) -> tuple[Path, Path]:
    exports_path.mkdir(parents=True, exist_ok=True)
    seed_default_budgets(conn, month)
    previous_month = _previous_month(month)
    current = transactions_for_month(conn, month)
    prior = transactions_for_month(conn, previous_month)
    summary = _summary_text(conn, month, previous_month, current, prior)
    markdown_path = exports_path / f"monthly_review_{month}.md"
    markdown_path.write_text(summary)
    conn.execute(
        """
        insert into monthly_reports(month, markdown_summary, created_at)
        values (?, ?, ?)
        on conflict(month) do update set
            markdown_summary = excluded.markdown_summary,
            created_at = excluded.created_at
        """,
        (month, summary, datetime.now(UTC).isoformat(timespec="seconds")),
    )
    conn.commit()

    workbook_path = exports_path / f"monthly_review_{month}.xlsx"
    _write_workbook(conn, month, previous_month, current, prior, summary, workbook_path)
    return workbook_path, markdown_path


def build_combined_report(
    conn: sqlite3.Connection,
    exports_path: Path,
    from_month: str | None = None,
    to_month: str | None = None,
) -> Path:
    exports_path.mkdir(parents=True, exist_ok=True)
    months = available_months(conn, from_month=from_month, to_month=to_month)
    if not months:
        raise ValueError("No transaction months are available for the requested range.")

    for month in months:
        seed_default_budgets(conn, month)

    workbook_path = exports_path / _combined_report_filename(months, from_month, to_month)
    monthly_rows = {month: transactions_for_month(conn, month) for month in months}
    monthly_category_totals = {
        month: category_totals(rows, conn, month)
        for month, rows in monthly_rows.items()
    }

    wb = Workbook()
    wb.remove(wb.active)
    _combined_overview_sheet(wb, conn, months, monthly_rows, monthly_category_totals)
    _combined_category_trend_sheet(wb, months, monthly_category_totals)
    _combined_focus_categories_sheet(wb, months, monthly_category_totals)
    _combined_top_merchants_sheet(wb, conn, months, monthly_rows)
    _combined_recurring_charges_sheet(wb, conn, months)
    _combined_category_drift_sheet(wb, conn, months)
    _combined_review_health_sheet(wb, conn, months, monthly_rows)
    _combined_itemized_allocations_sheet(wb, conn, months)
    _combined_subcategory_detail_sheet(wb, conn, months)
    _combined_raw_transactions_sheet(wb, conn, months)
    for ws in wb.worksheets:
        _style_header(ws)
        ws.freeze_panes = "A2"
    wb.save(workbook_path)
    return workbook_path


def available_months(
    conn: sqlite3.Connection,
    from_month: str | None = None,
    to_month: str | None = None,
) -> list[str]:
    months = [
        row["month"]
        for row in conn.execute(
            """
            select distinct substr(transaction_date, 1, 7) as month
            from transactions
            where transaction_date != ''
            order by month
            """
        )
    ]
    if from_month is not None:
        months = [month for month in months if month >= from_month]
    if to_month is not None:
        months = [month for month in months if month <= to_month]
    return months


def category_totals(
    rows: list[sqlite3.Row],
    conn: sqlite3.Connection | None = None,
    month: str | None = None,
) -> dict[str, float]:
    totals: dict[str, float] = defaultdict(float)
    allocations = itemized_category_allocations(conn, month) if conn is not None and month is not None else []
    split_transaction_ids = {allocation.transaction_id for allocation in allocations}
    for row in rows:
        if bool(row["exclude_from_spending"]):
            continue
        if int(row["id"]) in split_transaction_ids:
            continue
        category = row["category"] or "Needs Review"
        totals[category] += _spending_amount(row)
    for allocation in allocations:
        totals[allocation.category] += allocation.amount
    return {category: round(total, 2) for category, total in totals.items()}


def merchant_totals(
    rows: list[sqlite3.Row],
    conn: sqlite3.Connection | None = None,
    month: str | None = None,
) -> dict[tuple[str, str], dict[str, float]]:
    totals: dict[tuple[str, str], dict[str, float]] = {}
    allocations = itemized_category_allocations(conn, month) if conn is not None and month is not None else []
    split_transaction_ids = {allocation.transaction_id for allocation in allocations}
    for row in rows:
        if bool(row["exclude_from_spending"]):
            continue
        if int(row["id"]) in split_transaction_ids:
            continue
        key = (row["normalized_merchant"], row["category"] or "Needs Review")
        stats = totals.setdefault(key, {"count": 0, "total": 0.0})
        stats["count"] += 1
        stats["total"] += _spending_amount(row)
    for allocation in allocations:
        key = (allocation.merchant, allocation.category)
        stats = totals.setdefault(key, {"count": 0, "total": 0.0})
        stats["count"] += 1
        stats["total"] += allocation.amount
    return totals


def itemized_category_allocations(conn: sqlite3.Connection | None, month: str | None) -> list[ItemizedAllocation]:
    if conn is None or month is None:
        return []
    rows = list(
        conn.execute(
            """
            select
                t.id as transaction_id,
                t.normalized_merchant as merchant,
                t.amount as transaction_amount,
                t.exclude_from_spending as exclude_from_spending,
                i.item_price as item_price,
                i.quantity as quantity,
                i.item_total as item_total,
                i.item_category as item_category
            from transactions t
            join itemized_purchases i on i.transaction_id = t.id
            where substr(t.transaction_date, 1, 7) = ?
            order by t.id, i.id
            """,
            (month,),
        )
    )
    grouped: dict[int, list[sqlite3.Row]] = defaultdict(list)
    for row in rows:
        grouped[int(row["transaction_id"])].append(row)

    allocations: list[ItemizedAllocation] = []
    for transaction_id, items in grouped.items():
        if not items or bool(items[0]["exclude_from_spending"]):
            continue
        transaction_amount = _spending_amount(items[0], amount_column="transaction_amount")
        if transaction_amount <= 0:
            continue
        category_item_totals: dict[str, float] = defaultdict(float)
        for item in items:
            category = ITEM_TO_BUDGET.get(str(item["item_category"] or "Other"), "Other Discretionary")
            category_item_totals[category] += _item_total(item)
        item_total = sum(category_item_totals.values())
        if item_total <= 0:
            continue
        scaled = {
            category: (amount / item_total) * transaction_amount
            for category, amount in category_item_totals.items()
        }
        rounded = _round_allocations(scaled, transaction_amount)
        merchant = str(items[0]["merchant"])
        allocations.extend(
            ItemizedAllocation(
                transaction_id=transaction_id,
                merchant=merchant,
                category=category,
                amount=amount,
            )
            for category, amount in rounded.items()
            if amount > 0
        )
    return allocations


def itemized_subcategory_allocations(
    conn: sqlite3.Connection | None,
    month: str | None,
) -> list[ItemizedSubcategoryAllocation]:
    if conn is None or month is None:
        return []
    rows = list(
        conn.execute(
            """
            select
                t.id as transaction_id,
                t.normalized_merchant as merchant,
                t.amount as transaction_amount,
                t.exclude_from_spending as exclude_from_spending,
                i.item_price as item_price,
                i.quantity as quantity,
                i.item_total as item_total,
                i.item_category as item_category,
                i.item_subcategory as item_subcategory
            from transactions t
            join itemized_purchases i on i.transaction_id = t.id
            where substr(t.transaction_date, 1, 7) = ?
            order by t.id, i.id
            """,
            (month,),
        )
    )
    grouped: dict[int, list[sqlite3.Row]] = defaultdict(list)
    for row in rows:
        grouped[int(row["transaction_id"])].append(row)

    allocations: list[ItemizedSubcategoryAllocation] = []
    for transaction_id, items in grouped.items():
        if not items or bool(items[0]["exclude_from_spending"]):
            continue
        transaction_amount = _spending_amount(items[0], amount_column="transaction_amount")
        if transaction_amount <= 0:
            continue
        subcategory_item_totals: dict[tuple[str, str], float] = defaultdict(float)
        for item in items:
            item_category = str(item["item_category"] or "Other")
            category = ITEM_TO_BUDGET.get(item_category, "Other Discretionary")
            subcategory = _item_subcategory(item)
            subcategory_item_totals[(category, subcategory)] += _item_total(item)
        item_total = sum(subcategory_item_totals.values())
        if item_total <= 0:
            continue
        scaled = {
            key: (amount / item_total) * transaction_amount
            for key, amount in subcategory_item_totals.items()
        }
        rounded = _round_allocations({f"{category}\0{subcategory}": amount for (category, subcategory), amount in scaled.items()}, transaction_amount)
        merchant = str(items[0]["merchant"])
        for key, amount in rounded.items():
            if amount <= 0:
                continue
            category, subcategory = key.split("\0", 1)
            allocations.append(
                ItemizedSubcategoryAllocation(
                    transaction_id=transaction_id,
                    merchant=merchant,
                    category=category,
                    subcategory=subcategory,
                    amount=amount,
                )
            )
    return allocations


def _item_subcategory(item: sqlite3.Row) -> str:
    subcategory = str(item["item_subcategory"] or "").strip()
    if subcategory:
        return subcategory
    item_category = str(item["item_category"] or "Other")
    return {
        "Household Consumables": "Household Consumables",
        "Kids": "Kids Supplies",
        "Clothing": "Clothing",
        "Electronics": "Electronics Accessories",
        "Books / Media": "Books & Media",
        "Gifts": "Gifts",
        "Groceries": "Groceries",
    }.get(item_category, "General Discretionary")


def _round_allocations(amounts: dict[str, float], expected_total: float) -> dict[str, float]:
    rounded = {category: round(amount, 2) for category, amount in amounts.items()}
    delta = round(expected_total - sum(rounded.values()), 2)
    if rounded and abs(delta) >= 0.01:
        largest_category = max(rounded, key=lambda category: rounded[category])
        rounded[largest_category] = round(rounded[largest_category] + delta, 2)
    return rounded


def _item_total(item: sqlite3.Row) -> float:
    item_total = float(item["item_total"] or 0)
    if item_total > 0:
        return item_total
    return float(item["item_price"] or 0) * max(int(item["quantity"] or 1), 1)


def _spending_amount(row: sqlite3.Row, amount_column: str = "amount") -> float:
    amount = float(row[amount_column])
    if amount < 0:
        return abs(amount)
    return 0.0


def _summary_text(
    conn: sqlite3.Connection,
    month: str,
    previous_month: str,
    current: list[sqlite3.Row],
    prior: list[sqlite3.Row],
) -> str:
    totals = category_totals(current, conn, month)
    prior_totals = category_totals(prior, conn, previous_month)
    focus_lines = []
    for category in FOCUS_CATEGORIES:
        actual = totals.get(category, 0.0)
        prior_actual = prior_totals.get(category, 0.0)
        focus_lines.append(f"- {category}: ${actual:,.2f} ({actual - prior_actual:+,.2f} vs {previous_month})")
    review_count = sum(1 for row in current if bool(row["needs_review"]))
    total_spend = sum(totals.values())
    return "\n".join(
        [
            f"# Monthly Finance Review: {month}",
            "",
            f"Total tracked spending: ${total_spend:,.2f}",
            f"Needs review: {review_count} transaction(s)",
            "",
            "## Focus Categories",
            *focus_lines,
            "",
            "## Notes",
            "- Deterministic and human rules are applied before Ollama categorization.",
            "- Travel is tracked but not judged as overspending unless a budget is set.",
        ]
    )


def _write_workbook(
    conn: sqlite3.Connection,
    month: str,
    previous_month: str,
    current: list[sqlite3.Row],
    prior: list[sqlite3.Row],
    summary: str,
    path: Path,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _executive_summary_sheet(wb, summary)
    _scorecard_sheet(wb, conn, month, current, prior)
    _mom_sheet(wb, conn, month, previous_month, current, prior)
    _top_merchants_sheet(wb, conn, month, previous_month, current, prior)
    _subscriptions_sheet(wb, current)
    _itemized_allocations_sheet(wb, conn, month)
    _subcategory_detail_sheet(wb, conn, [month])
    _amazon_target_sheet(wb, conn, current)
    _needs_review_sheet(wb, current)
    _raw_transactions_sheet(wb, current)
    for ws in wb.worksheets:
        _style_header(ws)
    wb.save(path)


def _combined_report_filename(months: list[str], from_month: str | None, to_month: str | None) -> str:
    if from_month is None and to_month is None:
        return "combined_review_all.xlsx"
    start = from_month or months[0]
    end = to_month or months[-1]
    return f"combined_review_{start}_to_{end}.xlsx"


def _combined_overview_sheet(
    wb: Workbook,
    conn: sqlite3.Connection,
    months: list[str],
    monthly_rows: dict[str, list[sqlite3.Row]],
    monthly_category_totals: dict[str, dict[str, float]],
) -> None:
    ws = wb.create_sheet("Overview")
    total_spend = sum(sum(totals.values()) for totals in monthly_category_totals.values())
    total_transactions = sum(len(rows) for rows in monthly_rows.values())
    total_needs_review = sum(
        1
        for rows in monthly_rows.values()
        for row in rows
        if bool(row["needs_review"])
    )
    ws.append(["Metric", "Value"])
    ws.append(["Month Range", f"{months[0]} to {months[-1]}"])
    ws.append(["Months Included", len(months)])
    ws.append(["Total Spending", round(total_spend, 2)])
    ws.append(["Average Monthly Spending", round(total_spend / len(months), 2)])
    ws.append(["Transactions", total_transactions])
    ws.append(["Needs Review", total_needs_review])
    ws.append([])
    ws.append(["Month", "Transactions", "Total Spend", "Needs Review", "Itemized Matches", "Unmatched Itemized Orders"])
    for month in months:
        rows = monthly_rows[month]
        ws.append(
            [
                month,
                len(rows),
                round(sum(monthly_category_totals[month].values()), 2),
                sum(1 for row in rows if bool(row["needs_review"])),
                _itemized_match_count(conn, month),
                _unmatched_itemized_order_count(conn, month),
            ]
        )


def _combined_category_trend_sheet(
    wb: Workbook,
    months: list[str],
    monthly_category_totals: dict[str, dict[str, float]],
) -> None:
    ws = wb.create_sheet("Monthly Category Trend")
    ws.append(["Category", *months, "Total", "Average"])
    categories = sorted({category for totals in monthly_category_totals.values() for category in totals})
    for category in categories:
        values = [monthly_category_totals[month].get(category, 0.0) for month in months]
        ws.append([category, *values, round(sum(values), 2), round(sum(values) / len(months), 2)])


def _combined_focus_categories_sheet(
    wb: Workbook,
    months: list[str],
    monthly_category_totals: dict[str, dict[str, float]],
) -> None:
    ws = wb.create_sheet("Focus Categories")
    ws.append(["Month", *FOCUS_CATEGORIES, "Focus Total", "All Category Total"])
    for month in months:
        focus_values = [monthly_category_totals[month].get(category, 0.0) for category in FOCUS_CATEGORIES]
        ws.append(
            [
                month,
                *focus_values,
                round(sum(focus_values), 2),
                round(sum(monthly_category_totals[month].values()), 2),
            ]
        )


def _combined_top_merchants_sheet(
    wb: Workbook,
    conn: sqlite3.Connection,
    months: list[str],
    monthly_rows: dict[str, list[sqlite3.Row]],
) -> None:
    ws = wb.create_sheet("Top Merchants")
    ws.append(["Merchant", "Category", "Months Seen", "Transaction Count", "Total Spend", "Average Monthly Spend"])
    totals: dict[tuple[str, str], dict[str, object]] = {}
    for month in months:
        for key, stats in merchant_totals(monthly_rows[month], conn, month).items():
            aggregate = totals.setdefault(key, {"months": set(), "count": 0, "total": 0.0})
            aggregate["months"].add(month)
            aggregate["count"] = int(aggregate["count"]) + int(stats["count"])
            aggregate["total"] = float(aggregate["total"]) + float(stats["total"])
    for (merchant, category), stats in sorted(totals.items(), key=lambda item: float(item[1]["total"]), reverse=True):
        month_count = len(stats["months"])
        total = round(float(stats["total"]), 2)
        ws.append([merchant, category, month_count, stats["count"], total, round(total / len(months), 2)])


def _combined_recurring_charges_sheet(wb: Workbook, conn: sqlite3.Connection, months: list[str]) -> None:
    ws = wb.create_sheet("Recurring Charges")
    ws.append(["Merchant", "Amount", "Months Seen", "Count", "First Month", "Last Month", "Categories"])
    charges: dict[tuple[str, float], dict[str, object]] = {}
    for row in _transactions_between_months(conn, months):
        if bool(row["exclude_from_spending"]):
            continue
        amount = float(row["amount"])
        if amount >= 0:
            continue
        month = str(row["transaction_date"])[:7]
        key = (str(row["normalized_merchant"]), round(abs(amount), 2))
        stats = charges.setdefault(key, {"months": set(), "count": 0, "categories": set()})
        stats["months"].add(month)
        stats["count"] = int(stats["count"]) + 1
        stats["categories"].add(row["category"] or "Needs Review")
    for (merchant, amount), stats in sorted(charges.items(), key=lambda item: (-len(item[1]["months"]), item[0][0])):
        seen_months = sorted(stats["months"])
        if len(seen_months) < 2:
            continue
        ws.append(
            [
                merchant,
                amount,
                len(seen_months),
                stats["count"],
                seen_months[0],
                seen_months[-1],
                " | ".join(sorted(stats["categories"])),
            ]
        )


def _combined_category_drift_sheet(wb: Workbook, conn: sqlite3.Connection, months: list[str]) -> None:
    ws = wb.create_sheet("Category Drift")
    ws.append(["Merchant", "Categories", "Months Seen", "Transaction Count", "Total Spend"])
    merchants: dict[str, dict[str, object]] = {}
    for row in _transactions_between_months(conn, months):
        if bool(row["exclude_from_spending"]):
            continue
        merchant = str(row["normalized_merchant"])
        stats = merchants.setdefault(merchant, {"categories": set(), "months": set(), "count": 0, "total": 0.0})
        stats["categories"].add(row["category"] or "Needs Review")
        stats["months"].add(str(row["transaction_date"])[:7])
        stats["count"] = int(stats["count"]) + 1
        stats["total"] = float(stats["total"]) + _spending_amount(row)
    for merchant, stats in sorted(merchants.items()):
        categories = sorted(stats["categories"])
        if len(categories) < 2:
            continue
        ws.append(
            [
                merchant,
                " | ".join(categories),
                len(stats["months"]),
                stats["count"],
                round(float(stats["total"]), 2),
            ]
        )


def _combined_review_health_sheet(
    wb: Workbook,
    conn: sqlite3.Connection,
    months: list[str],
    monthly_rows: dict[str, list[sqlite3.Row]],
) -> None:
    ws = wb.create_sheet("Review Health")
    ws.append(
        [
            "Month",
            "Transactions",
            "Needs Review",
            "Review %",
            "Itemized Matches",
            "Unmatched Itemized Orders",
            "Unmatched Department Store Transactions",
        ]
    )
    for month in months:
        rows = monthly_rows[month]
        needs_review = sum(1 for row in rows if bool(row["needs_review"]))
        review_ratio = needs_review / len(rows) if rows else 0.0
        ws.append(
            [
                month,
                len(rows),
                needs_review,
                round(review_ratio, 4),
                _itemized_match_count(conn, month),
                _unmatched_itemized_order_count(conn, month),
                _unmatched_department_store_transaction_count(conn, month),
            ]
        )


def _combined_itemized_allocations_sheet(wb: Workbook, conn: sqlite3.Connection, months: list[str]) -> None:
    ws = wb.create_sheet("Itemized Allocations")
    ws.append(["Month", "Transaction ID", "Merchant", "Category", "Allocated Amount"])
    for month in months:
        for allocation in itemized_category_allocations(conn, month):
            ws.append([month, allocation.transaction_id, allocation.merchant, allocation.category, allocation.amount])


def _combined_subcategory_detail_sheet(wb: Workbook, conn: sqlite3.Connection, months: list[str]) -> None:
    _subcategory_detail_sheet(wb, conn, months)


def _subcategory_detail_sheet(wb: Workbook, conn: sqlite3.Connection, months: list[str]) -> None:
    ws = wb.create_sheet("Subcategory Detail")
    ws.append(["Month", "Budget Category", "Subcategory", "Merchant", "Transaction Count", "Allocated Amount"])
    aggregates: dict[tuple[str, str, str, str], dict[str, object]] = {}
    for month in months:
        split_transaction_ids = {
            allocation.transaction_id
            for allocation in itemized_subcategory_allocations(conn, month)
        }
        for allocation in itemized_subcategory_allocations(conn, month):
            if allocation.category not in {"Household & Kids", "Other Discretionary", "Groceries"}:
                continue
            key = (month, allocation.category, allocation.subcategory, allocation.merchant)
            stats = aggregates.setdefault(key, {"transactions": set(), "amount": 0.0})
            stats["transactions"].add(allocation.transaction_id)
            stats["amount"] = float(stats["amount"]) + allocation.amount
        for row in transactions_for_month(conn, month):
            if bool(row["exclude_from_spending"]) or int(row["id"]) in split_transaction_ids:
                continue
            category = row["category"] or "Needs Review"
            subcategory = str(row["subcategory"] or "").strip()
            if category not in {"Household & Kids", "Other Discretionary", "Groceries"} or not subcategory:
                continue
            key = (month, category, subcategory, row["normalized_merchant"])
            stats = aggregates.setdefault(key, {"transactions": set(), "amount": 0.0})
            stats["transactions"].add(int(row["id"]))
            stats["amount"] = float(stats["amount"]) + _spending_amount(row)
    for (month, category, subcategory, merchant), stats in sorted(
        aggregates.items(),
        key=lambda item: (item[0][0], item[0][1], -float(item[1]["amount"])),
    ):
        ws.append(
            [
                month,
                category,
                subcategory,
                merchant,
                len(stats["transactions"]),
                round(float(stats["amount"]), 2),
            ]
        )


def _combined_raw_transactions_sheet(wb: Workbook, conn: sqlite3.Connection, months: list[str]) -> None:
    ws = wb.create_sheet("Raw Combined Transactions")
    ws.append(
        [
            "ID",
            "Date",
            "Month",
            "Merchant",
            "Description",
            "Source Category",
            "Amount",
            "Category",
            "Subcategory",
            "Confidence",
            "Source",
            "Needs Review",
            "Exclude From Spending",
        ]
    )
    for row in _transactions_between_months(conn, months):
        ws.append(
            [
                row["id"],
                row["transaction_date"],
                str(row["transaction_date"])[:7],
                row["normalized_merchant"],
                row["raw_description"],
                row["source_category"],
                row["amount"],
                row["category"],
                row["subcategory"],
                row["confidence"],
                row["categorization_source"],
                bool(row["needs_review"]),
                bool(row["exclude_from_spending"]),
            ]
        )


def _transactions_between_months(conn: sqlite3.Connection, months: list[str]) -> list[sqlite3.Row]:
    return list(
        conn.execute(
            """
            select *
            from transactions
            where substr(transaction_date, 1, 7) between ? and ?
            order by transaction_date, id
            """,
            (months[0], months[-1]),
        )
    )


def _itemized_match_count(conn: sqlite3.Connection, month: str) -> int:
    return int(
        conn.execute(
            """
            select count(distinct i.transaction_id)
            from itemized_purchases i
            join transactions t on t.id = i.transaction_id
            where substr(t.transaction_date, 1, 7) = ?
            """,
            (month,),
        ).fetchone()[0]
        or 0
    )


def _unmatched_itemized_order_count(conn: sqlite3.Connection, month: str) -> int:
    return int(
        conn.execute(
            """
            select count(*)
            from (
                select merchant, order_id, order_date
                from itemized_purchases
                where transaction_id is null
                    and substr(order_date, 1, 7) = ?
                group by merchant, order_id, order_date
            )
            """,
            (month,),
        ).fetchone()[0]
        or 0
    )


def _unmatched_department_store_transaction_count(conn: sqlite3.Connection, month: str) -> int:
    return int(
        conn.execute(
            """
            select count(*)
            from transactions
            where substr(transaction_date, 1, 7) = ?
                and normalized_merchant in ('Amazon', 'Target', 'Walmart', 'Costco')
                and amount < 0
                and exclude_from_spending = 0
                and id not in (
                    select transaction_id
                    from itemized_purchases
                    where transaction_id is not null
                )
            """,
            (month,),
        ).fetchone()[0]
        or 0
    )


def _executive_summary_sheet(wb: Workbook, summary: str) -> None:
    ws = wb.create_sheet("Executive Summary")
    for index, line in enumerate(summary.splitlines(), start=1):
        ws.cell(index, 1, line)
    ws.column_dimensions["A"].width = 100


def _scorecard_sheet(wb: Workbook, conn: sqlite3.Connection, month: str, current, prior) -> None:
    ws = wb.create_sheet("Category Scorecard")
    ws.append(["Category", "Budget", "Actual", "Difference", "Prior Month", "MoM Change"])
    budgets = {
        row["category"]: row["budget_amount"]
        for row in conn.execute("select category, budget_amount from monthly_budgets where month = ?", (month,))
    }
    totals = category_totals(current, conn, month)
    prior_totals = category_totals(prior, conn, _previous_month(month))
    for category in FOCUS_CATEGORIES:
        actual = totals.get(category, 0.0)
        budget = budgets.get(category, 0.0)
        prior_actual = prior_totals.get(category, 0.0)
        ws.append([category, budget, actual, budget - actual, prior_actual, actual - prior_actual])


def _mom_sheet(wb: Workbook, conn: sqlite3.Connection, month: str, previous_month: str, current, prior) -> None:
    ws = wb.create_sheet("Month-over-Month")
    ws.append(["Category", "Current", previous_month, "Change"])
    totals = category_totals(current, conn, month)
    prior_totals = category_totals(prior, conn, previous_month)
    for category in sorted(set(totals) | set(prior_totals)):
        ws.append([category, totals.get(category, 0.0), prior_totals.get(category, 0.0), totals.get(category, 0.0) - prior_totals.get(category, 0.0)])


def _top_merchants_sheet(wb: Workbook, conn: sqlite3.Connection, month: str, previous_month: str, current, prior) -> None:
    ws = wb.create_sheet("Top Merchants")
    ws.append(["Merchant", "Category", "Transaction Count", "Total Spend", "Average Transaction", "Prior Month Total", "Change"])
    totals = merchant_totals(current, conn, month)
    prior_totals = merchant_totals(prior, conn, previous_month)
    for (merchant, category), stats in sorted(totals.items(), key=lambda item: item[1]["total"], reverse=True):
        prior_total = prior_totals.get((merchant, category), {}).get("total", 0.0)
        average = stats["total"] / stats["count"] if stats["count"] else 0.0
        ws.append([merchant, category, stats["count"], round(stats["total"], 2), round(average, 2), round(prior_total, 2), round(stats["total"] - prior_total, 2)])


def _subscriptions_sheet(wb: Workbook, current) -> None:
    ws = wb.create_sheet("Subscriptions")
    ws.append(["Date", "Merchant", "Description", "Amount", "Confidence"])
    for row in current:
        if row["category"] == "Subscriptions":
            ws.append([row["transaction_date"], row["normalized_merchant"], row["raw_description"], row["amount"], row["confidence"]])


def _amazon_target_sheet(wb: Workbook, conn: sqlite3.Connection, current) -> None:
    ws = wb.create_sheet("Amazon Target Detail")
    ws.append(["Type", "Date/Order", "Merchant", "Description/Item", "Amount", "Category", "Subcategory", "Confidence"])
    for row in current:
        if row["normalized_merchant"] in {"Amazon", "Target"}:
            ws.append([
                "Transaction",
                row["transaction_date"],
                row["normalized_merchant"],
                row["raw_description"],
                row["amount"],
                row["category"],
                row["subcategory"],
                row["confidence"],
            ])
    for item in itemized_purchases(conn):
        ws.append([
            "Item",
            item["order_id"],
            item["merchant"],
            item["item_title"],
            item["item_price"],
            item["item_category"],
            item["item_subcategory"],
            item["confidence"],
        ])


def _itemized_allocations_sheet(wb: Workbook, conn: sqlite3.Connection, month: str) -> None:
    ws = wb.create_sheet("Itemized Allocations")
    ws.append(["Transaction ID", "Merchant", "Category", "Allocated Amount"])
    for allocation in itemized_category_allocations(conn, month):
        ws.append([allocation.transaction_id, allocation.merchant, allocation.category, allocation.amount])


def _needs_review_sheet(wb: Workbook, current) -> None:
    ws = wb.create_sheet("Needs Review")
    ws.append(["Date", "Merchant", "Description", "Amount", "Suggested Category", "Suggested Subcategory", "Confidence", "Reason"])
    for row in current:
        if bool(row["needs_review"]):
            ws.append([row["transaction_date"], row["normalized_merchant"], row["raw_description"], row["amount"], row["category"], row["subcategory"], row["confidence"], row["notes"]])


def _raw_transactions_sheet(wb: Workbook, current) -> None:
    ws = wb.create_sheet("Raw Transactions")
    ws.append(["ID", "Date", "Merchant", "Description", "Source Category", "Amount", "Category", "Subcategory", "Confidence", "Source", "Exclude From Spending"])
    for row in current:
        ws.append([row["id"], row["transaction_date"], row["normalized_merchant"], row["raw_description"], row["source_category"], row["amount"], row["category"], row["subcategory"], row["confidence"], row["categorization_source"], bool(row["exclude_from_spending"])])


def _style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="D9EAD3")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 48)


def _previous_month(month: str) -> str:
    year, month_number = [int(part) for part in month.split("-")]
    if month_number == 1:
        return f"{year - 1}-12"
    return f"{year}-{month_number - 1:02d}"
